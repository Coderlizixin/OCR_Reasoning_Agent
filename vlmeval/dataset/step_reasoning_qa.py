"""
Step Reasoning QA Dataset
==========================
VLMEvalKit dataset classes:
  - Step_Reasoning_QA: evaluate models on step-level reasoning QA pairs
  - OCR_Reasoning_StepQA: OCR_Reasoning with step QA as additional context
"""
from ..smp import *
from .image_base import ImageBaseDataset
from .image_vqa import OCR_Reasoning
from collections import defaultdict


# ============================================================
# Helper: load step QA pairs for OCR_Reasoning_StepQA
# ============================================================

_STEP_QA_CACHE = None  # module-level cache


def _load_step_qa_pairs():
    """Load step QA pairs from step_qa_pairs.xlsx. Returns dict[question_id] -> list of steps."""
    global _STEP_QA_CACHE
    if _STEP_QA_CACHE is not None:
        return _STEP_QA_CACHE

    import openpyxl
    import json

    candidates = [
        r'T:\PyCharmProject\OCR-Reasoning\OCR-Reasoning-main\outputs\step_qa\step_qa_pairs.xlsx',
        'outputs/step_qa/step_qa_pairs.xlsx',
    ]
    path = None
    for c in candidates:
        if osp.exists(c):
            path = c
            break
    if path is None:
        print('[OCR_Reasoning_StepQA] step_qa_pairs.xlsx not found, step QA augmentation disabled')
        _STEP_QA_CACHE = {}
        return _STEP_QA_CACHE

    wb = openpyxl.load_workbook(path)
    ws = wb.active

    qa_map = defaultdict(list)
    for row in range(2, ws.max_row + 1):
        qid = int(ws.cell(row, 1).value)
        qa_map[qid].append({
            'step_num': ws.cell(row, 5).value,
            'total_steps': ws.cell(row, 6).value,
            'qa_question': str(ws.cell(row, 9).value),
            'qa_answer': str(ws.cell(row, 10).value),
            'qa_type': str(ws.cell(row, 11).value),
            'qa_options': str(ws.cell(row, 12).value),
        })

    _STEP_QA_CACHE = dict(qa_map)
    return _STEP_QA_CACHE


# ============================================================
# Dataset: Step_Reasoning_QA (standalone step-level evaluation)
# ============================================================

class Step_Reasoning_QA(ImageBaseDataset):
    """Step-level Reasoning QA Benchmark.

    Evaluates whether a model can answer decomposed reasoning steps correctly.
    Success criterion: ALL steps must be correct (question-level accuracy).
    Also tracks per-step accuracy for detailed analysis.
    """

    TYPE = 'VQA'
    DATASET_URL = {
        'Step_Reasoning_QA': 'outputs/step_qa/step_reasoning_qa.tsv'
    }
    DATASET_MD5 = {'Step_Reasoning_QA': ''}  # local file, no MD5 check

    # It returns a DataFrame
    @classmethod
    def evaluate(self, eval_file, **judge_kwargs):
        """Evaluate step-level QA predictions using GPT judge.

        Same judge flow as OCR_Reasoning: build_judge → extract answers → compare.
        """
        from .utils.step_reasoning_qa import evaluate_step_qa
        from .utils.judge_util import build_judge

        # Build judge model (same way OCR_Reasoning does)
        model = build_judge(max_tokens=1024, **judge_kwargs)

        result = evaluate_step_qa(eval_file, judge_model=model)
        return result

    def build_prompt(self, line):
        """Build prompt: image + step questions + format instruction.

        Uses absolute image paths directly from the 'image' column (bypasses
        dump_image since our images are already files, not base64).
        """
        if isinstance(line, int):
            line = self.data.iloc[line]

        # Image is already an absolute file path (> 64 chars)
        tgt_path = str(line['image'])

        question = line['question']
        language = line.get('language', 'English')
        is_cn = 'Chinese' in str(language)

        # Add format instruction matching OCR_Reasoning style
        if is_cn:
            question += "\n通过一步一步的推理解决以上子问题。按顺序输出每个子问题的答案：子问题1: [答案], 子问题2: [答案], ..."
        else:
            question += "\nSolve the above sub-questions through step-by-step reasoning. Output your answer for each sub-question in order: Sub-question 1: [answer], Sub-question 2: [answer], ..."

        msgs = []
        if isinstance(tgt_path, list):
            msgs.extend([dict(type='image', value=p) for p in tgt_path])
        else:
            msgs = [dict(type='image', value=tgt_path)]
        msgs.append(dict(type='text', value=question))
        return msgs

    # Override load_data to handle local TSV
    def load_data(self, dataset):
        """Load from local TSV file instead of URL download."""
        rel_path = self.DATASET_URL.get(dataset, 'outputs/step_qa/step_reasoning_qa.tsv')

        # Try multiple locations
        candidates = [
            rel_path,                                          # relative to cwd
            osp.join(LMUDataRoot(), rel_path),                 # under LMUData
            osp.join(LMUDataRoot(), dataset + '.tsv'),         # under LMUData root
            r'T:\PyCharmProject\OCR-Reasoning\OCR-Reasoning-main\outputs\step_qa\step_reasoning_qa.tsv',  # absolute
        ]
        data_path = None
        for c in candidates:
            if osp.exists(c):
                data_path = c
                break

        if data_path is None:
            raise FileNotFoundError(
                f"Dataset TSV not found. Tried: {candidates}. "
                f"Run build_step_reasoning_qa_tsv.py first."
            )
        return load(data_path)


# ============================================================
# Dataset: OCR_Reasoning_StepQA
# ============================================================
# OCR_Reasoning benchmark with step-level QA pairs injected as auxiliary context.
# For a question with n steps, only the first n-1 step QA pairs are included
# (the nth step contains the final answer, which would leak the solution).
#
# Usage:
#   python run.py --data OCR_Reasoning_StepQA --model Qwen3.5-9B-API
#
# Everything else (data loading, evaluation, judge) is identical to OCR_Reasoning.

class OCR_Reasoning_StepQA(OCR_Reasoning):
    """OCR_Reasoning with step-level QA pairs as additional hints.

    Inherits evaluate() from OCR_Reasoning — uses the same GPT judge.
    Only build_prompt() is modified to inject the first n-1 step QA pairs.
    """

    DATASET_URL = {
        'OCR_Reasoning_StepQA': OCR_Reasoning.DATASET_URL['OCR_Reasoning']
    }
    DATASET_MD5 = {'OCR_Reasoning_StepQA': OCR_Reasoning.DATASET_MD5['OCR_Reasoning']}

    def __init__(self, dataset='OCR_Reasoning_StepQA', **kwargs):
        super().__init__(dataset=dataset, **kwargs)
        self.step_qa_map = _load_step_qa_pairs()
        n_with_qa = sum(1 for steps in self.step_qa_map.values() if len(steps) > 1)
        print(f'[OCR_Reasoning_StepQA] Loaded step QA for {len(self.step_qa_map)} questions, '
              f'{n_with_qa} have >=2 steps (can inject n-1 hints)')

    def build_prompt(self, line):
        """Build prompt with step QA hints injected before the format instruction.

        Original OCR_Reasoning prompt:
          {image}
          {question}
          通过一步一步的推理解决这个复杂的问题。最后的回答的组成应该是: '{format}'.

        Augmented prompt:
          {image}
          {question}

          以下中间推理步骤及其答案可供参考：
          步骤1: {qa_question_1}
          答案: {qa_answer_1}
          步骤2: {qa_question_2}
          答案: {qa_answer_2}
          ... (first n-1 steps only, excluding the final answer step)

          通过一步一步的推理解决这个复杂的问题。最后的回答的组成应该是: '{format}'.
        """
        if isinstance(line, int):
            line = self.data.iloc[line]

        # Get base messages (image + question) from parent
        msgs = super(OCR_Reasoning, self).build_prompt(line)
        # Note: we call OCR_Reasoning's parent (ImageBaseDataset.build_prompt),
        # NOT OCR_Reasoning.build_prompt, because we'll add the instruction ourselves.

        qid = int(line['index']) if 'index' in line else None
        steps = self.step_qa_map.get(qid, [])
        total_steps = len(steps)

        # Inject first n-1 step QA pairs (exclude last step = final answer)
        hint_steps = steps[:-1] if total_steps > 1 else []

        # Build the augmented text
        text_item = None
        for item in msgs:
            if item['type'] == 'text':
                text_item = item
                break

        if text_item is None:
            return msgs

        is_cn = 'Chinese' in str(line.get('language', ''))

        if hint_steps:
            if is_cn:
                parts = ["\n\n以下中间推理步骤及其答案可供参考："]
                for s in hint_steps:
                    sn = s['step_num']
                    sq = s['qa_question']
                    sa = s['qa_answer']
                    parts.append(f"步骤{sn}: {sq}\n答案: {sa}")
                parts.append("")
            else:
                parts = ["\n\nThe following intermediate reasoning steps and their answers "
                         "are provided for reference:"]
                for s in hint_steps:
                    sn = s['step_num']
                    sq = s['qa_question']
                    sa = s['qa_answer']
                    parts.append(f"Step {sn}: {sq}\nAnswer: {sa}")
                parts.append("")
            hint_text = "\n".join(parts)
            text_item['value'] += hint_text

        # Add the original OCR_Reasoning format instruction
        fmt = line.get('format', '')
        if is_cn:
            text_item['value'] += f"\n通过一步一步的推理解决这个复杂的问题。最后的回答的组成应该是: '{fmt}'."
        else:
            text_item['value'] += f"\nSolve the complex problem through step-by-step reasoning. "
            text_item['value'] += f"The composition of the final answer should be: '{fmt}'."

        return msgs


# ============================================================
# Dataset: Depended_Step_Reasoning_QA
# ============================================================
# Step-level QA with dependency-aware prompting.
# Only the first 40 questions (id <= 269) from the OCR-Reasoning random subset.
#
# For steps with dependencies, the prompt includes the dependent step Q&A
# as a reference block. Multiple-choice answers show the full option text.
# Question-level scoring: all steps must be correct.
#
# Usage:
#   python run.py --data Depended_Step_Reasoning_QA --model Qwen3.5-9B-API --judge qwen-plus

class Depended_Step_Reasoning_QA(Step_Reasoning_QA):
    """Depended Step Reasoning QA — with dependency-aware prompting.

    Inherits evaluate() from Step_Reasoning_QA (GPT-judge + per-step checking).
    The prompt already contains the dependency reference block (built into the
    TSV question field). Only overrides the data source.
    """

    DATASET_URL = {
        'Depended_Step_Reasoning_QA':
            'outputs/step_qa_depended/step_reasoning_qa.tsv'
    }
    DATASET_MD5 = {'Depended_Step_Reasoning_QA': ''}

    def load_data(self, dataset):
        """Load from the depended TSV file."""
        candidates = [
            osp.abspath('outputs/step_qa_depended/step_reasoning_qa.tsv'),
            r'T:\PyCharmProject\OCR-Reasoning\OCR-Reasoning-main\outputs\step_qa_depended\step_reasoning_qa.tsv',
        ]
        data_path = None
        for c in candidates:
            if osp.exists(c):
                data_path = c
                break
        if data_path is None:
            raise FileNotFoundError(
                f"Depended TSV not found. Run build_depended_tsv.py first."
            )
        return load(data_path)
