"""
Build Step Reasoning QA Dataset TSV
====================================
Reads step_qa_pairs.xlsx and aggregates QA pairs per question_id
into a single dataset row for VLMEvalKit evaluation.

Output: step_reasoning_qa.tsv — each row = one original question with all step QAs
"""
import openpyxl
import json
import os
import pandas as pd
from collections import defaultdict

INPUT_EXCEL = "outputs/step_qa/step_qa_pairs.xlsx"
OUTPUT_TSV = "outputs/step_qa/step_reasoning_qa.tsv"
IMAGE_DIR = "outputs/step_qa/images"


def build_dataset_tsv():
    wb = openpyxl.load_workbook(INPUT_EXCEL)
    ws = wb.active

    # Group by question_id
    questions = defaultdict(list)
    for row in range(2, ws.max_row + 1):
        q_id = int(ws.cell(row, 1).value)
        questions[q_id].append({
            "step_num": ws.cell(row, 5).value,
            "total_steps": ws.cell(row, 6).value,
            "step_text": str(ws.cell(row, 7).value),
            "qa_question": str(ws.cell(row, 9).value),
            "qa_answer": str(ws.cell(row, 10).value),
            "qa_type": str(ws.cell(row, 11).value),
            "qa_options": str(ws.cell(row, 12).value),
            "qa_cot": str(ws.cell(row, 13).value),
        })

    # Also get original question info from the Excel
    wb_in = openpyxl.load_workbook(
        "outputs/Qwen3.5-9B-OCRAgentV5/T20260730_G533ef742/Qwen3.5-9B-OCRAgentV5_OCR_Reasoning.xlsx"
    )
    ws_in = wb_in.active
    orig_info = {}
    for row in range(2, ws_in.max_row + 1):
        q_id = int(ws_in.cell(row, 1).value)
        orig_info[q_id] = {
            "question": str(ws_in.cell(row, 2).value),
            "language": str(ws_in.cell(row, 9).value),
            "task": str(ws_in.cell(row, 7).value) if ws_in.cell(row, 7).value else "",
        }

    # Build the prompt format for each question
    rows = []
    for q_id in sorted(questions.keys()):
        steps = questions[q_id]
        total = steps[0]["total_steps"]
        lang = orig_info.get(q_id, {}).get("language", "English")
        is_cn = "Chinese" in lang

        # Build multi-step QA prompt
        if is_cn:
            prompt_parts = ["请根据图片内容，逐步回答以下子问题：\n"]
            for s in steps:
                sn = s["step_num"]
                sq = s["qa_question"]
                st = s["qa_type"]
                opts = s["qa_options"]

                if st == "multiple_choice" and opts and opts != "[]":
                    try:
                        options = json.loads(opts)
                        opt_text = "\n".join(f"{o['label']}. {o['text']}" for o in options)
                        prompt_parts.append(f"步骤{sn}: {sq}\n选项:\n{opt_text}")
                    except:
                        prompt_parts.append(f"步骤{sn}: {sq}")
                else:
                    prompt_parts.append(f"步骤{sn}: {sq}")

            prompt_parts.append(
                f"\n请按照以下格式输出每个步骤的答案：\n"
                f"步骤1: [你的答案]\n步骤2: [你的答案]\n...\n"
                f"注意：选择题只需输出选项字母（如A/B/C/D），简答题直接输出答案值。"
            )
        else:
            prompt_parts = ["Please answer the following sub-questions step by step based on the image:\n"]
            for s in steps:
                sn = s["step_num"]
                sq = s["qa_question"]
                st = s["qa_type"]
                opts = s["qa_options"]

                if st == "multiple_choice" and opts and opts != "[]":
                    try:
                        options = json.loads(opts)
                        opt_text = "\n".join(f"{o['label']}. {o['text']}" for o in options)
                        prompt_parts.append(f"Step {sn}: {sq}\nOptions:\n{opt_text}")
                    except:
                        prompt_parts.append(f"Step {sn}: {sq}")
                else:
                    prompt_parts.append(f"Step {sn}: {sq}")

            prompt_parts.append(
                f"\nPlease output your answer for each step in the following format:\n"
                f"Step 1: [your answer]\nStep 2: [your answer]\n...\n"
                f"Note: For multiple choice, output only the option letter (e.g., A/B/C/D). "
                f"For short answer, output the value directly."
            )

        full_prompt = "\n".join(prompt_parts)

        # Build the ground truth answer (JSON with all step answers)
        gt_answers = {}
        for s in steps:
            gt_answers[f"step{s['step_num']}"] = {
                "answer": s["qa_answer"],
                "type": s["qa_type"],
            }

        # Image path — MUST be absolute and > 64 chars (VLMEvalKit constraint)
        img_path = os.path.abspath(os.path.join(IMAGE_DIR, f"{q_id}.jpg"))
        if not os.path.exists(img_path):
            img_path = os.path.abspath(os.path.join(
                r"C:\Users\19639\LMUData\images\OCR_Reasoning", f"{q_id}.jpg"
            ))
        if not os.path.exists(img_path):
            print(f"Warning: image not found for question_id={q_id}")
            img_path = ""

        rows.append({
            "index": q_id,
            "question": full_prompt,
            "answer": json.dumps(gt_answers, ensure_ascii=False),
            "image": img_path,
            "language": lang,
            "task": orig_info.get(q_id, {}).get("task", ""),
            "step_count": total,
            "orig_question": orig_info.get(q_id, {}).get("question", ""),
        })

    df = pd.DataFrame(rows)
    df.to_csv(OUTPUT_TSV, sep="\t", index=False)
    print(f"Dataset saved: {OUTPUT_TSV}")
    print(f"  Questions: {len(df)}")
    print(f"  Total steps: {df['step_count'].sum()}")
    print(f"  Chinese: {(df['language'] == 'Chinese').sum()}")
    print(f"  English: {(df['language'] == 'English').sum()}")

    # Print a sample
    print("\n--- Sample (first question) ---")
    row = df.iloc[0]
    print(f"index={row['index']}, language={row['language']}, steps={row['step_count']}")
    print(f"Prompt preview:\n{row['question'][:500]}...")
    print(f"Answer: {row['answer'][:200]}...")

    return df


if __name__ == "__main__":
    build_dataset_tsv()
