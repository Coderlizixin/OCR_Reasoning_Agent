import openpyxl, re, sys, io, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('outputs/Qwen3.5-9B-OCRAgentV5/T20260730_G533ef742/Qwen3.5-9B-OCRAgentV5_OCR_Reasoning.xlsx')
ws = wb.active

def parse_cn_steps(text):
    parts = re.split(r'(?:第[一二三四五六七八九十百零\d]+步)[：:]', text)
    if parts and '最终的回答是' in parts[-1]:
        parts[-1] = parts[-1].split('最终的回答是')[0].strip()
    return [s.strip() for s in parts if s.strip()]

def parse_en_steps(text):
    parts = re.split(r'(?:step\s*\d+)\s*:?\s*', text, flags=re.IGNORECASE)
    if parts and 'the final answer is' in parts[-1].lower():
        parts[-1] = re.split(r'the final answer is', parts[-1], flags=re.IGNORECASE)[0].strip()
    return [s.strip() for s in parts if s.strip()]

def assess_qa_feasibility(step, step_idx, total_steps):
    issues = []

    item_separators = re.findall(r'[、,，]', step)
    bullet_items = len(re.findall(r'(?:^|\n)\s*\*', step))
    numbers_listed = len(re.findall(r'\d+\.\d+%|\d+\.\d+元|\d+元|\d+\.\d+万|\d+万|\d+\.\d+亿', step))

    total_items = len(item_separators) + bullet_items

    if total_items > 12:
        issues.append(('enumeration_too_long', '列举了约%d个项目，做成选择题选项过多，做成问答题答案太长难以验证' % total_items))

    if re.search(r'\\frac|\\sqrt|\\int|\\sum|\\prod|\\lim|\\begin|\\end|\\overrightarrow|\\angle|\\triangle|\\sin|\\cos|\\tan|\\log', step):
        issues.append(('complex_math', '包含LaTeX公式，是数学推导过程，难以简单验证'))

    if re.search(r'识别.*类别|识别.*种类|识别.*类型', step) and total_items > 4:
        issues.append(('multi_category_extraction', '识别并列举多个类别/种类，答案是一个长列表'))

    if re.search(r'^(分析|判断|综合|综上|总结|根据.*可知|由.*可知|Analyze|Deduce|State|Note that)', step) and not re.search(r'\d+', step):
        issues.append(('descriptive_only', '纯粹描述性/判断性步骤，没有具体可验证的数值答案'))

    if re.search(r'经过的站点|站点有|站。|stations|bus.*route|stop', step) and total_items > 10:
        issues.append(('route_enumeration', '列举大量站点名称，难以逐一验证'))

    if re.search(r'最低价格|最低价|最高价格|最高价|cheapest|most expensive|lowest price|highest price', step) and total_items > 8:
        issues.append(('min_max_from_many', '从大量数值中找最大/最小值，列举过程冗长'))

    if (step_idx == total_steps - 1 or step_idx == total_steps) and re.search(r'^(正确|错误|The answer is|故.*正确|故.*错误|综上.*选|因此.*选)', step):
        issues.append(('answer_statement', '最后一步仅陈述答案选项，本身不是独立的推理步骤'))

    if re.search(r'图.*显示|图.*表示|图表.*展示|chart.*shows|graph.*shows|image.*shows|根据.*图表|from the chart', step) and total_items > 6:
        issues.append(('chart_description', '描述图表内容的步骤，包含多个数据点'))

    if re.search(r'价格.*\d+.*\d+.*\d+.*\d+', step) and numbers_listed > 6:
        issues.append(('multi_price_compare', '列举并比较大量价格数据'))

    if not issues:
        has_single_result = bool(re.search(r'[=＝]|结果为|结果为|答案是|得到|得出|is|result|answer', step))
        has_single_number = len(re.findall(r'\d+\.?\d*', step)) <= 5
        is_short = len(step) < 200

        if has_single_result and has_single_number and is_short:
            return ('easy', '单步计算或提取，结果是一个明确可验证的数值', 'short_answer')
        elif re.search(r'判断|选择|choose|select|pick|which one|which option|correct|incorrect', step, re.IGNORECASE):
            return ('easy', '判断或选择题型，可天然转化为选择题', 'multiple_choice')
        elif total_items <= 5 and has_single_number:
            return ('easy', '少量项目的列举或计算，结果明确', 'short_answer_or_mc')
        else:
            return ('medium', '步骤有一定复杂度但可转化为问答题', 'short_answer')

    if any(i[0] in ('enumeration_too_long', 'route_enumeration', 'complex_math') for i in issues):
        return ('hard', '; '.join(i[1] for i in issues), 'not_easily_convertible')
    elif any(i[0] in ('answer_statement', 'multi_category_extraction', 'min_max_from_many', 'chart_description', 'multi_price_compare') for i in issues):
        return ('hard', '; '.join(i[1] for i in issues), 'partially_convertible')
    else:
        return ('medium', '; '.join(i[1] for i in issues), 'adaptable')


# ===== MAIN ANALYSIS =====
from collections import Counter, defaultdict

feasibility_counter = Counter()
problem_details = defaultdict(list)

cn_easy = cn_medium = cn_hard = 0
en_easy = en_medium = en_hard = 0

for row in range(2, ws.max_row + 1):
    r = str(ws.cell(row, 4).value)
    lang = str(ws.cell(row, 9).value)
    q = str(ws.cell(row, 2).value)[:150]

    if 'chinese' in lang.lower():
        steps = parse_cn_steps(r)
        lang_type = 'cn'
    else:
        steps = parse_en_steps(r)
        lang_type = 'en'

    for i, step in enumerate(steps):
        feasibility, reason, suggested = assess_qa_feasibility(step, i+1, len(steps))
        feasibility_counter[feasibility] += 1

        if lang_type == 'cn':
            if feasibility == 'easy': cn_easy += 1
            elif feasibility == 'medium': cn_medium += 1
            elif feasibility == 'hard': cn_hard += 1
        else:
            if feasibility == 'easy': en_easy += 1
            elif feasibility == 'medium': en_medium += 1
            elif feasibility == 'hard': en_hard += 1

        if feasibility == 'hard':
            problem_details[reason.split(';')[0].strip()].append({
                'row': row, 'lang': lang_type, 'step_num': i+1, 'total': len(steps),
                'text_preview': step[:200], 'question': q[:100]
            })

print('=' * 70)
print('      推理步骤QA转化可行性分析报告')
print('=' * 70)
print()
print('总题数: %d (中文77题, 英文123题)' % (ws.max_row - 1))
print('总步骤数: %d (中文%d步, 英文%d步)' % (feasibility_counter.total(), cn_easy+cn_medium+cn_hard, en_easy+en_medium+en_hard))
print()
print('--- 按语言分布的可行性 ---')
print('%-12s %8s %8s %8s' % ('', 'Easy', 'Medium', 'Hard'))
print('%-12s %8d %8d %8d' % ('中文(CN)', cn_easy, cn_medium, cn_hard))
print('%-12s %8d %8d %8d' % ('英文(EN)', en_easy, en_medium, en_hard))
print('%-12s %8d %8d %8d' % ('合计', cn_easy+en_easy, cn_medium+en_medium, cn_hard+en_hard))
print()

print('--- 可行性分布 ---')
for k in ['easy', 'medium', 'hard']:
    count = feasibility_counter.get(k, 0)
    pct = count / feasibility_counter.total() * 100
    bar = '#' * int(pct / 2)
    print('  %s: %d步 (%.1f%%) %s' % (k, count, pct, bar))

print()
print('=' * 70)
print('      主要问题类别及典型例子')
print('=' * 70)

problem_names = {
    'enumeration_too_long': '列举项目过多(>12项)',
    'complex_math': '复杂数学公式/推导(LaTeX)',
    'multi_category_extraction': '多类别识别提取',
    'descriptive_only': '纯描述性/判断性步骤',
    'route_enumeration': '路线/站点列举',
    'min_max_from_many': '从大量数值中找最值',
    'answer_statement': '答案陈述步骤(非推理)',
    'chart_description': '图表内容描述',
    'multi_price_compare': '多产品价格比较',
}

for key, name in problem_names.items():
    cases = problem_details.get(key, [])
    if not cases:
        continue
    print('\n### %s (%d个步骤)' % (name, len(cases)))
    for c in cases[:2]:
        print('  题目(row %d): %s' % (c['row'], c['question']))
        print('  步骤%d/%d [%s]: %s' % (c['step_num'], c['total'], c['lang'], c['text_preview'][:180]))
        print()

print()
print('=' * 70)
print('      每道题的步骤可行性概览')
print('=' * 70)

# Per-question summary with highlighting for hard steps
questions_with_hard_steps = []
for row in range(2, ws.max_row + 1):
    r = str(ws.cell(row, 4).value)
    lang = str(ws.cell(row, 9).value)
    q = str(ws.cell(row, 2).value)[:100]

    if 'chinese' in lang.lower():
        steps = parse_cn_steps(r)
        lang_type = 'CN'
    else:
        steps = parse_en_steps(r)
        lang_type = 'EN'

    assessments = []
    for i, step in enumerate(steps):
        f, _, _ = assess_qa_feasibility(step, i+1, len(steps))
        assessments.append(f[0].upper())

    step_status = ' '.join('S%d:%s' % (s+1, a) for s, a in enumerate(assessments))
    hard_count = sum(1 for a in assessments if a == 'H')
    flag = ' *** HASSUE' if hard_count > 0 else ''
    print('Row%d [%s %d步%s]: %s' % (row, lang_type, len(steps), flag, q))
    print('  Steps: %s' % step_status)
    if hard_count > 0:
        questions_with_hard_steps.append((row, lang_type, len(steps), hard_count, q[:80]))

print()
print('=' * 70)
print('      存在难以转化步骤的题目汇总 (%d题)' % len(questions_with_hard_steps))
print('=' * 70)
for row, lt, total, hc, q in questions_with_hard_steps:
    print('  Row%d [%s] %d/%d步困难: %s' % (row, lt, hc, total, q))
