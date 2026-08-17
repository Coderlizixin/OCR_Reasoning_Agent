"""Verify the generated step QA pairs output."""
import openpyxl, json, os, sys
from collections import Counter

def verify_output(excel_path: str, image_dir: str):
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    headers = [str(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    n_rows = ws.max_row - 1

    print(f"Output Excel: {excel_path}")
    print(f"Total QA pairs: {n_rows}")
    print(f"Headers: {headers}")
    print()

    # Statistics
    q_ids = set()
    gen_methods = Counter()
    qa_types = Counter()
    languages = Counter()
    empty_answers = 0
    empty_questions = 0
    mc_with_no_options = 0

    for row in range(2, ws.max_row + 1):
        q_id = ws.cell(row, 1).value
        q_ids.add(q_id)
        gen_methods[str(ws.cell(row, 15).value)] += 1
        qa_types[str(ws.cell(row, 11).value)] += 1
        languages[str(ws.cell(row, 3).value)] += 1

        qa_question = str(ws.cell(row, 9).value) if ws.cell(row, 9).value else ""
        qa_answer = str(ws.cell(row, 10).value) if ws.cell(row, 10).value else ""
        qa_type = str(ws.cell(row, 11).value) if ws.cell(row, 11).value else ""
        qa_options = ws.cell(row, 12).value

        if not qa_question.strip():
            empty_questions += 1
        if not qa_answer.strip() or qa_answer == "[需要人工标注]":
            empty_answers += 1
        if qa_type == "multiple_choice" and (not qa_options or qa_options == "[]"):
            mc_with_no_options += 1

    print(f"Unique questions: {len(q_ids)}")
    print(f"Languages: {dict(languages)}")
    print(f"QA Types: {dict(qa_types)}")
    print(f"Generation methods: {dict(gen_methods)}")
    print()
    print(f"Empty questions: {empty_questions}")
    print(f"Empty/placeholder answers: {empty_answers}")
    print(f"Multiple choice without options: {mc_with_no_options}")
    print()

    # Check images
    if os.path.exists(image_dir):
        images = os.listdir(image_dir)
        print(f"Images in output dir: {len(images)}")
        print(f"  Sample: {images[:5]}")

    print()
    print("--- Sample QA pairs (first 10) ---")
    for row in range(2, min(12, ws.max_row + 1)):
        q_id = ws.cell(row, 1).value
        step_num = ws.cell(row, 5).value
        total = ws.cell(row, 6).value
        qa_q = str(ws.cell(row, 9).value)[:120] if ws.cell(row, 9).value else "(empty)"
        qa_a = str(ws.cell(row, 10).value)[:80] if ws.cell(row, 10).value else "(empty)"
        qa_t = ws.cell(row, 11).value
        gen = ws.cell(row, 15).value
        print(f"  [{q_id}] Step{step_num}/{total} [{qa_t}] [{gen}]")
        print(f"    Q: {qa_q}")
        print(f"    A: {qa_a}")
        print()

if __name__ == "__main__":
    excel_path = sys.argv[1] if len(sys.argv) > 1 else "outputs/step_qa/step_qa_pairs.xlsx"
    image_dir = sys.argv[2] if len(sys.argv) > 2 else "outputs/step_qa/images"
    verify_output(excel_path, image_dir)
