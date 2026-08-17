"""
Build Depended Step Reasoning QA TSV
=====================================
Reads outputs/step_qa_depended/step_qa_pairs.xlsx and produces a TSV for
VLMEvalKit's Depended_Step_Reasoning_QA dataset.

For steps with dependencies, the prompt includes the dependent Q&A as a
reference block. Multiple-choice dependency answers show the full option text.
"""
import openpyxl
import json
import os
import shutil
from collections import defaultdict
import pandas as pd

INPUT_EXCEL = "outputs/step_qa_depended/step_qa_pairs.xlsx"
OUTPUT_DIR = "outputs/step_qa_depended"
OUTPUT_TSV = f"{OUTPUT_DIR}/step_reasoning_qa.tsv"
OUTPUT_IMAGES = f"{OUTPUT_DIR}/images"
SOURCE_IMAGES = "outputs/step_qa/images"
FALLBACK_IMAGES = r"C:\Users\19639\LMUData\images\OCR_Reasoning"

# Only include questions with id <= this value (first ~40 questions)
MAX_QUESTION_ID = 269

# Original question info for language/task
ORIG_EXCEL = "outputs/Qwen3.5-9B-OCRAgentV5/T20260730_G533ef742/Qwen3.5-9B-OCRAgentV5_OCR_Reasoning.xlsx"


def get_option_text(letter: str, options_str: str) -> str:
    """For a multiple-choice answer letter, get the full option text.

    Handles both formats:
      - ["A. xxx", "B. xxx", ...]  (standardized string format)
      - [{"label": "A", "text": "xxx"}, ...]  (legacy dict format)
    """
    if not options_str or options_str == "[]" or not letter:
        return letter
    try:
        options = json.loads(options_str)
        letter_upper = letter.strip().upper()

        for opt in options:
            if isinstance(opt, str):
                # Format: "A. xxx"
                if opt.strip().startswith(letter_upper + "."):
                    return opt.strip()
            elif isinstance(opt, dict):
                # Legacy format: {"label": "A", "text": "xxx"}
                if opt.get("label", "").upper() == letter_upper:
                    return f"{letter_upper}. {opt.get('text', '')}"

        return letter  # fallback
    except (json.JSONDecodeError, TypeError):
        return letter


def build_tsv():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    os.makedirs(OUTPUT_IMAGES, exist_ok=True)

    # Load original question info
    wb_orig = openpyxl.load_workbook(ORIG_EXCEL)
    ws_orig = wb_orig.active
    orig_info = {}
    for row in range(2, ws_orig.max_row + 1):
        qid = int(ws_orig.cell(row, 1).value)
        orig_info[qid] = {
            "question": str(ws_orig.cell(row, 2).value),
            "language": str(ws_orig.cell(row, 9).value),
            "task": str(ws_orig.cell(row, 7).value) if ws_orig.cell(row, 7).value else "",
        }

    # Load depended step QA pairs
    wb = openpyxl.load_workbook(INPUT_EXCEL)
    ws = wb.active

    qa_groups = defaultdict(list)
    for row in range(2, ws.max_row + 1):
        qid = int(ws.cell(row, 1).value)
        if qid > MAX_QUESTION_ID:
            continue  # only first ~40 questions
        qa_groups[qid].append({
            "question_id": qid,
            "question": str(ws.cell(row, 2).value),
            "language": str(ws.cell(row, 3).value),
            "step_num": ws.cell(row, 5).value,
            "total_steps": ws.cell(row, 6).value,
            "qa_question": str(ws.cell(row, 9).value),
            "qa_answer": str(ws.cell(row, 10).value),
            "qa_type": str(ws.cell(row, 11).value),
            "qa_options": str(ws.cell(row, 12).value),
            "depends_on": str(ws.cell(row, 16).value) if ws.cell(row, 16).value else "",
        })

    print(f"Loaded {len(qa_groups)} questions, {sum(len(v) for v in qa_groups.values())} steps")
    print(f"Question IDs: {sorted(qa_groups.keys())}")

    ts_rows = []

    for qid in sorted(qa_groups.keys()):
        steps = sorted(qa_groups[qid], key=lambda x: x["step_num"])
        info = orig_info.get(qid, {})
        lang = info.get("language", "English")
        is_cn = "Chinese" in lang
        total = len(steps)

        # Build lookup: step_num -> answer (with full option text for MC)
        answers_by_step = {}
        questions_by_step = {}
        for s in steps:
            sn = s["step_num"]
            questions_by_step[sn] = s["qa_question"]
            if s["qa_type"] == "multiple_choice":
                full_text = get_option_text(s["qa_answer"], s["qa_options"])
                answers_by_step[sn] = full_text
            else:
                answers_by_step[sn] = s["qa_answer"]

        # Collect steps that are depended on by any later step
        all_depended = set()
        for s in steps:
            dep_str = s["depends_on"]
            if dep_str:
                for dsn in dep_str.split(","):
                    dsn = dsn.strip()
                    if dsn:
                        all_depended.add(int(dsn))

        # Build reference block
        ref_parts = []
        if all_depended:
            if is_cn:
                ref_parts.append("以下中间步骤及其答案可供后续步骤参考：")
            else:
                ref_parts.append("The following intermediate steps and their "
                                 "answers are provided for reference in later steps:")
            for dsn in sorted(all_depended):
                dq = questions_by_step.get(dsn, "")
                da = answers_by_step.get(dsn, "")
                if is_cn:
                    ref_parts.append(f"步骤{dsn}: {dq}")
                    ref_parts.append(f"答案: {da}")
                else:
                    ref_parts.append(f"Step {dsn}: {dq}")
                    ref_parts.append(f"Answer: {da}")
            ref_parts.append("")
        ref_block = "\n".join(ref_parts)

        # Build step questions section
        if is_cn:
            parts = ["请根据图片内容，逐步回答以下子问题：\n"]
            for s in steps:
                sn = s["step_num"]
                sq = s["qa_question"]
                st = s["qa_type"]
                opts = s["qa_options"]
                if st == "multiple_choice" and opts and opts != "[]":
                    try:
                        options = json.loads(opts)
                        opt_lines = []
                        for o in options:
                            if isinstance(o, str):
                                opt_lines.append(o)
                            else:
                                opt_lines.append(f"{o['label']}. {o['text']}")
                        opt_text = "\n".join(opt_lines)
                        parts.append(f"子问题{sn}: {sq}\n选项:\n{opt_text}")
                    except:
                        parts.append(f"子问题{sn}: {sq}")
                else:
                    parts.append(f"子问题{sn}: {sq}")
            parts.append(
                "\n按顺序输出每个子问题的答案。答案格式："
                "\n子问题1: [你的答案]\n子问题2: [你的答案]\n...\n"
                "注意：选择题只输出选项字母（A/B/C/D），简答题直接输出数值或结果。"
            )
        else:
            parts = ["Please answer the following sub-questions step by step "
                     "based on the image:\n"]
            for s in steps:
                sn = s["step_num"]
                sq = s["qa_question"]
                st = s["qa_type"]
                opts = s["qa_options"]
                if st == "multiple_choice" and opts and opts != "[]":
                    try:
                        options = json.loads(opts)
                        opt_lines = []
                        for o in options:
                            if isinstance(o, str):
                                opt_lines.append(o)
                            else:
                                opt_lines.append(f"{o['label']}. {o['text']}")
                        opt_text = "\n".join(opt_lines)
                        parts.append(f"Sub-question {sn}: {sq}\nOptions:\n{opt_text}")
                    except:
                        parts.append(f"Sub-question {sn}: {sq}")
                else:
                    parts.append(f"Sub-question {sn}: {sq}")
            parts.append(
                "\nOutput your answer for each sub-question in order. "
                "Answer format:\n"
                "Sub-question 1: [your answer]\n"
                "Sub-question 2: [your answer]\n...\n"
                "Note: For multiple-choice, output only the option letter (A/B/C/D). "
                "For short-answer, output the value or result directly."
            )

        question_text = "\n".join(parts)
        if ref_block:
            question_text = ref_block + "\n" + question_text

        # GT answers
        gt_answers = {}
        for s in steps:
            gt_answers[f"step{s['step_num']}"] = {
                "answer": s["qa_answer"],
                "type": s["qa_type"],
            }

        # Image path
        img_path = os.path.abspath(os.path.join(OUTPUT_IMAGES, f"{qid}.jpg"))
        if not os.path.exists(img_path):
            src = os.path.join(SOURCE_IMAGES, f"{qid}.jpg")
            if not os.path.exists(src):
                src = os.path.join(FALLBACK_IMAGES, f"{qid}.jpg")
            if os.path.exists(src):
                shutil.copy2(src, img_path)

        ts_rows.append({
            "index": qid,
            "question": question_text,
            "answer": json.dumps(gt_answers, ensure_ascii=False),
            "image": img_path,
            "language": lang,
            "task": info.get("task", ""),
            "step_count": total,
            "orig_question": info.get("question", ""),
        })

    df = pd.DataFrame(ts_rows)
    df.to_csv(OUTPUT_TSV, sep="\t", index=False)

    # Verify image paths
    for i, row in df.iterrows():
        p = str(row['image'])
        if len(p) <= 64:
            print(f"WARNING: short path for qid={row['index']}: {p} (len={len(p)})")
        if not os.path.exists(p):
            print(f"WARNING: image not found for qid={row['index']}: {p}")

    print(f"\nTSV saved: {OUTPUT_TSV}")
    print(f"  Questions: {len(df)}")
    print(f"  Total steps: {df['step_count'].sum()}")
    print(f"  Images: {len(os.listdir(OUTPUT_IMAGES)) if os.path.exists(OUTPUT_IMAGES) else 0}")

    # Show Q13 prompt
    row = df[df['index'] == 13].iloc[0]
    print(f"\n=== Q13 prompt preview ===")
    print(row['question'][:600])


if __name__ == "__main__":
    build_tsv()
